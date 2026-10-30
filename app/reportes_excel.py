"""Genera el Excel descargable de la página de Reportes (ver la ruta
/dashboard/reportes/excel en main.py).

Junta cuatro cosas en un solo archivo, cada una con su propia hoja:

1. Resumen diario — a partir de con qué trámite empezó cada solicitud (creado_en, que
   nunca cambia) y en qué quedó ahora mismo (estado actual): cuántas llegaron con
   documentos completos, cuántas requirieron asesor, cuántas siguen pendientes.
2. Pendientes por gestionar — el backlog ACTUAL (no limitado al rango de fechas
   elegido): cada solicitud que todavía no se gestionó, sin importar cuándo se creó,
   con su estado, de quién es, y cuántos días lleva esperando (ver
   `solicitudes.abiertas`).
3. Gestionados por día + detalle — usando actualizado_en, que para una solicitud ya
   gestionada sí refleja exactamente el momento de la gestión (ver
   `solicitudes.gestionadas_entre`).
4. Preguntas frecuentes — le pasamos a Claude los mensajes de los pacientes en el
   rango (sin teléfono ni nombre) y le pedimos que los agrupe por tema.
"""
import io
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime, timezone

from anthropic import Anthropic
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.config import settings
from app.solicitudes import NOMBRE_SERVICIO, Solicitud

_client = Anthropic(api_key=settings.anthropic_api_key)

NAVY = "FF041E42"
TEAL = "FF10616F"
WHITE = "FFFFFFFF"
BEIGE = "FFF3EFE8"

ETIQUETA_ESTADO = {
    "documentos_completos": "Documentos completos",
    "agendada": "Agendada",
    "requiere_humano": "Requiere asesor",
    "atendida": "Atendida",
    "pendiente_documentos": "Pendiente de documentos",
    "orden_vencida": "Orden vencida / incompleta",
    "no_interesado": "No interesado",
}


# ---------------------------------------------------------------------------
# Preguntas frecuentes — agrupación por tema con Claude
# ---------------------------------------------------------------------------

AGRUPAR_TEMAS_TOOL = {
    "name": "agrupar_temas",
    "description": (
        "Agrupa mensajes de pacientes de WhatsApp en temas recurrentes, con cuántos "
        "mensajes corresponden a cada tema y un par de ejemplos textuales."
    ),
    "input_schema": {
        "type": "object",
        "properties": {
            "temas": {
                "type": "array",
                "items": {
                    "type": "object",
                    "properties": {
                        "tema": {
                            "type": "string",
                            "description": "Nombre corto del tema, ej. 'Horarios de atención', 'Documentos faltantes'",
                        },
                        "cantidad": {
                            "type": "integer",
                            "description": "Cuántos de los mensajes dados corresponden a este tema",
                        },
                        "ejemplos": {
                            "type": "array",
                            "items": {"type": "string"},
                            "description": "1 a 3 mensajes textuales de ejemplo, copiados tal cual (sin corregirles nada)",
                        },
                    },
                    "required": ["tema", "cantidad", "ejemplos"],
                },
            }
        },
        "required": ["temas"],
    },
}


@dataclass
class TemaPregunta:
    tema: str
    cantidad: int
    ejemplos: list[str]


def agrupar_preguntas_frecuentes(mensajes: list[str]) -> list[TemaPregunta]:
    """Le pide a Claude que agrupe los mensajes de los pacientes (ya sin teléfono ni
    nombre) en temas recurrentes, ordenados del más al menos frecuente. Si no hay
    mensajes en el rango, no llama a la API — devuelve una lista vacía."""
    if not mensajes:
        return []

    # Hasta 400 mensajes por llamada — de sobra para ver los temas más comunes sin un
    # prompt gigante. Si hay más, se toma una muestra repartida a lo largo de todo el
    # rango (no solo los primeros que llegaron) para no perderse temas que aparecieron
    # más adelante.
    muestra = mensajes
    if len(mensajes) > 400:
        paso = len(mensajes) / 400
        muestra = [mensajes[int(i * paso)] for i in range(400)]

    texto_mensajes = "\n".join(f"- {m}" for m in muestra)

    response = _client.messages.create(
        model=settings.anthropic_model,
        max_tokens=2048,
        system=(
            "Vas a recibir mensajes que pacientes le escribieron por WhatsApp a Aurora, "
            "el asistente de trámites de audífonos de Compensar / Widex Colombia S.A.S. "
            "Agrúpalos en TEMAS recurrentes — no un tema por mensaje: junta los que "
            "preguntan lo mismo aunque estén escritos distinto (ej. 'a q hora abren' y "
            "'cual es el horario' van al mismo tema). Ordena los temas del más "
            "frecuente al menos frecuente. Ignora saludos sueltos ('hola', 'buenos "
            "días'), confirmaciones de datos (cédula, nombre, sí/no) y respuestas de "
            "menú (escribir solo un número) — quédate con preguntas, dudas y quejas "
            "reales sobre el trámite."
        ),
        tools=[AGRUPAR_TEMAS_TOOL],
        tool_choice={"type": "tool", "name": "agrupar_temas"},
        messages=[{"role": "user", "content": f"Estos son los mensajes:\n{texto_mensajes}"}],
    )

    for block in response.content:
        if block.type == "tool_use" and block.name == "agrupar_temas":
            temas = block.input.get("temas", [])
            return [
                TemaPregunta(
                    tema=str(t.get("tema", "")),
                    cantidad=int(t.get("cantidad") or 0),
                    ejemplos=[str(e) for e in (t.get("ejemplos") or [])][:3],
                )
                for t in temas
            ]
    return []


# ---------------------------------------------------------------------------
# Construcción del archivo .xlsx
# ---------------------------------------------------------------------------

def _dia(iso: str) -> str:
    """Solo la fecha (AAAA-MM-DD) de un timestamp ISO guardado en Supabase — para
    agrupar por día sin que la hora exacta estorbe. '(sin fecha)' si el dato vino
    vacío o raro, en vez de reventar el reporte completo por una fila mala."""
    if not iso:
        return "(sin fecha)"
    try:
        return datetime.fromisoformat(iso.replace("Z", "+00:00")).date().isoformat()
    except ValueError:
        return "(sin fecha)"


def _encabezado(ws, fila: int, columnas: list[str], color: str = NAVY) -> None:
    relleno = PatternFill(start_color=color, end_color=color, fill_type="solid")
    for col_idx, titulo in enumerate(columnas, start=1):
        celda = ws.cell(row=fila, column=col_idx, value=titulo)
        celda.font = Font(bold=True, color=WHITE)
        celda.fill = relleno
        celda.alignment = Alignment(vertical="center", wrap_text=True)


def _autoancho(ws, columnas: list[str], anchos: list[int]) -> None:
    for col_idx, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = ancho


def construir_excel(
    desde: date,
    hasta: date,
    cohorte: list[Solicitud],
    gestionadas: list[Solicitud],
    temas: list[TemaPregunta],
    abiertas: list[Solicitud] | None = None,
) -> bytes:
    wb = Workbook()

    # --- Hoja 1: Resumen diario (según cuándo EMPEZÓ cada solicitud) ----------------
    ws1 = wb.active
    ws1.title = "Resumen diario"
    ws1.cell(row=1, column=1, value=f"Aurora — Resumen del {desde.isoformat()} al {hasta.isoformat()}").font = Font(
        bold=True, size=13, color=NAVY
    )
    ws1.cell(
        row=2,
        column=1,
        value="Por fecha en que se inició cada trámite — a partir de ahí, en qué quedó cada uno hoy.",
    ).font = Font(italic=True, color="FF666666")

    columnas1 = ["Fecha", "Trámites iniciados", "Documentos completos", "Requirieron asesor",
                 "Pendientes de documentos", "Orden vencida / incompleta"]
    _encabezado(ws1, 4, columnas1)

    por_dia: dict[str, dict[str, int]] = defaultdict(
        lambda: {"total": 0, "completos": 0, "asesor": 0, "pendientes": 0, "orden_vencida": 0}
    )
    for s in cohorte:
        fila = por_dia[_dia(s.creado_en)]
        fila["total"] += 1
        if s.estado in ("documentos_completos", "agendada"):
            fila["completos"] += 1
        elif s.estado in ("requiere_humano", "atendida"):
            fila["asesor"] += 1
        elif s.estado == "pendiente_documentos":
            fila["pendientes"] += 1
        elif s.estado == "orden_vencida":
            fila["orden_vencida"] += 1

    fila_actual = 5
    for dia in sorted(por_dia):
        datos = por_dia[dia]
        ws1.append([dia, datos["total"], datos["completos"], datos["asesor"], datos["pendientes"], datos["orden_vencida"]])
        fila_actual += 1
    if not por_dia:
        ws1.cell(row=fila_actual, column=1, value="No hubo trámites iniciados en este rango de fechas.")

    _autoancho(ws1, columnas1, [14, 18, 20, 18, 22, 24])

    # --- Hoja 2: Pendientes por gestionar (backlog ACTUAL, no limitado al rango) ----
    # A diferencia de "Resumen diario" (que solo cuenta lo que se INICIÓ en el rango
    # elegido), esta hoja muestra, ahora mismo, cada solicitud que todavía no se
    # gestionó — sin importar cuándo empezó — para que no se pierda de vista una que
    # lleva semanas esperando solo porque quedó fuera del rango de fechas del reporte.
    ws_pend = wb.create_sheet("Pendientes por gestionar", 1)
    ws_pend.cell(row=1, column=1, value="Aurora — Pendientes por gestionar (al momento de generar este Excel)").font = Font(
        bold=True, size=13, color=NAVY
    )
    ws_pend.cell(
        row=2,
        column=1,
        value="Documentos completos, requieren asesor, pendientes de documentos y orden vencida — todo lo que falta por gestionar ahora mismo, no solo lo del rango de fechas elegido.",
    ).font = Font(italic=True, color="FF666666")

    columnas_pend = ["Estado", "Nombre", "Cédula", "Teléfono", "Trámite", "Fecha de creación", "Días esperando"]
    _encabezado(ws_pend, 4, columnas_pend)

    ahora = datetime.now(timezone.utc)

    def _dias_esperando(iso: str) -> str:
        if not iso:
            return ""
        try:
            desde_dt = datetime.fromisoformat(iso.replace("Z", "+00:00"))
        except ValueError:
            return ""
        return str(max(0, (ahora - desde_dt).days))

    fila_actual = 5
    # Las más urgentes (más días esperando) primero.
    for s in sorted(abiertas or [], key=lambda s: s.actualizado_en or "", reverse=False):
        ws_pend.append([
            ETIQUETA_ESTADO.get(s.estado, s.estado),
            s.nombre,
            s.cedula,
            s.telefono,
            NOMBRE_SERVICIO.get(s.servicio, s.servicio),
            _dia(s.creado_en),
            _dias_esperando(s.actualizado_en),
        ])
        fila_actual += 1
    if not abiertas:
        ws_pend.cell(row=fila_actual, column=1, value="No hay ninguna solicitud pendiente por gestionar ahora mismo.")
    _autoancho(ws_pend, columnas_pend, [20, 26, 14, 16, 34, 18, 16])
    ws_pend.freeze_panes = "A5"

    # --- Hoja 3: Gestionados por día (según cuándo se GESTIONÓ) ---------------------
    ws2 = wb.create_sheet("Gestionados por día")
    columnas2 = ["Fecha de gestión", "Cuántos se gestionaron"]
    _encabezado(ws2, 1, columnas2, color=TEAL)
    gestion_por_dia: dict[str, int] = defaultdict(int)
    for s in gestionadas:
        gestion_por_dia[_dia(s.actualizado_en)] += 1
    fila_actual = 2
    for dia in sorted(gestion_por_dia):
        ws2.cell(row=fila_actual, column=1, value=dia)
        ws2.cell(row=fila_actual, column=2, value=gestion_por_dia[dia])
        fila_actual += 1
    if not gestion_por_dia:
        ws2.cell(row=fila_actual, column=1, value="No se gestionó ningún caso en este rango de fechas.")

    # --- Mismo hoja: cuántos se gestionaron de cada trámite, en este rango ----------
    fila_tramite = fila_actual + 2
    columnas2b = ["Trámite", "Gestionados en el rango", "Atendida", "Agendada"]
    _encabezado(ws2, fila_tramite, columnas2b, color=TEAL)

    por_tramite: dict[str, dict[str, int]] = defaultdict(lambda: {"total": 0, "atendida": 0, "agendada": 0})
    for s in gestionadas:
        fila = por_tramite[s.servicio]
        fila["total"] += 1
        if s.estado == "atendida":
            fila["atendida"] += 1
        elif s.estado == "agendada":
            fila["agendada"] += 1

    fila_actual2 = fila_tramite + 1
    # Mismo orden que el menú de Aurora (ver solicitudes.SERVICIOS), no alfabético, para
    # que se lea igual de fácil que el resto del reporte — y solo los trámites que sí
    # tuvieron al menos una gestión en el rango, para no llenar la hoja de ceros.
    for codigo in NOMBRE_SERVICIO:
        if codigo not in por_tramite:
            continue
        datos = por_tramite[codigo]
        ws2.cell(row=fila_actual2, column=1, value=NOMBRE_SERVICIO.get(codigo, codigo))
        ws2.cell(row=fila_actual2, column=2, value=datos["total"])
        ws2.cell(row=fila_actual2, column=3, value=datos["atendida"])
        ws2.cell(row=fila_actual2, column=4, value=datos["agendada"])
        fila_actual2 += 1
    if not por_tramite:
        ws2.cell(row=fila_actual2, column=1, value="No se gestionó ningún caso en este rango de fechas.")

    # Las dos tablas comparten columnas 1-2 (la de abajo además usa 3-4), así que los
    # anchos se fijan una sola vez al final, pensando en ambas.
    _autoancho(ws2, ["", "", "", ""], [34, 24, 14, 14])

    # --- Hoja 4: Detalle de gestionados ---------------------------------------------
    ws3 = wb.create_sheet("Detalle de gestionados")
    columnas3 = ["Fecha de gestión", "Nombre", "Cédula", "Teléfono", "Trámite", "Resultado", "Gestionado por", "Observaciones"]
    _encabezado(ws3, 1, columnas3)
    fila_actual = 2
    for s in sorted(gestionadas, key=lambda s: s.actualizado_en, reverse=True):
        ws3.append([
            _dia(s.actualizado_en),
            s.nombre,
            s.cedula,
            s.telefono,
            NOMBRE_SERVICIO.get(s.servicio, s.servicio),
            ETIQUETA_ESTADO.get(s.estado, s.estado),
            s.atendido_por,
            s.observaciones,
        ])
        fila_actual += 1
    if not gestionadas:
        ws3.cell(row=fila_actual, column=1, value="No se gestionó ningún caso en este rango de fechas.")
    _autoancho(ws3, columnas3, [16, 26, 14, 16, 34, 18, 26, 44])
    for fila in ws3.iter_rows(min_row=2, max_col=8):
        fila[7].alignment = Alignment(wrap_text=True, vertical="top")

    # --- Hoja 5: Preguntas frecuentes -------------------------------------------------
    ws4 = wb.create_sheet("Preguntas frecuentes")
    columnas4 = ["Tema", "Cantidad de mensajes", "Ejemplo 1", "Ejemplo 2", "Ejemplo 3"]
    _encabezado(ws4, 1, columnas4, color=TEAL)
    fila_actual = 2
    for t in temas:
        ejemplos = t.ejemplos + ["", "", ""]
        ws4.append([t.tema, t.cantidad, ejemplos[0], ejemplos[1], ejemplos[2]])
        fila_actual += 1
    if not temas:
        ws4.cell(row=fila_actual, column=1, value="No hubo suficientes mensajes de pacientes en este rango para agrupar por tema.")
    _autoancho(ws4, columnas4, [28, 20, 40, 40, 40])
    for fila in ws4.iter_rows(min_row=2, max_col=5):
        for celda in fila[2:]:
            celda.alignment = Alignment(wrap_text=True, vertical="top")

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
